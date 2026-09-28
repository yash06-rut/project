import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

template_path = r"f:\final\E-commerce1\E-commerce1\Internship_UDP_Training_Report_Template for reporting 4.docx"
output_path = r"f:\final\E-commerce1\E-commerce1\Filled_Internship_UDP_Training_Report_4.docx"

doc = docx.Document(template_path)

def format_run(run, font_size=10, bold=False, italic=False, color_rgb=(34, 34, 34)):
    run.font.name = "Segoe UI"
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

# 1. Fill Student Information Table (Table 0)
if len(doc.tables) > 0:
    t0 = doc.tables[0]
    try:
        t0.cell(0, 1).text = "Krish"
        t0.cell(1, 1).text = "22SSO2CA005"
        t0.cell(2, 1).text = "Computer Engineering"
        t0.cell(3, 1).text = "Semester 6"
        t0.cell(4, 1).text = "+91 9876543210"
        t0.cell(5, 1).text = "krish@example.com"
    except Exception:
        pass

# 2. Fill Internship / UDP Details Table (Table 1)
if len(doc.tables) > 1:
    t1 = doc.tables[1]
    try:
        t1.cell(0, 1).text = "User Defined Project (UDP)"
        t1.cell(1, 1).text = "P P Savani University"
        t1.cell(2, 1).text = "Full Stack Web Development / E-Commerce"
        t1.cell(3, 1).text = "Industry Guide"
        t1.cell(4, 1).text = "Institute Mentor"
        t1.cell(5, 1).text = "01-06-2026"
        t1.cell(6, 1).text = "25-07-2026"
    except Exception:
        pass

for i in range(min(2, len(doc.tables))):
    table = doc.tables[i]
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    format_run(r, font_size=10)

# Process paragraphs
for p in doc.paragraphs:
    text = p.text.strip()

    # Project Title
    if text.startswith("Project Title:"):
        p.text = "Project Title: Full-Stack E-Commerce Web Application (Aura)"
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            format_run(r, font_size=10.5, bold=True, color_rgb=(17, 17, 17))
            
    elif "________________________________________" in text:
        p.text = "" 

    elif "(Mention development / research work completed)" in text:
        p.text = (
            "Successfully designed and implemented a full-stack responsive e-commerce web application. "
            "Built the frontend user interface using HTML/CSS/JS and the backend REST API using Node.js, Express, and MongoDB. "
            "Integrated secure user authentication with JWT and payment processing functionality with Razorpay."
        )
        for r in p.runs: format_run(r, font_size=10)

    elif "(Attach architecture, block diagram, flowchart, UML diagram, etc.)" in text:
        p.text = (
            "System Architecture Overview:\n"
            "• Frontend Client: HTML5, CSS3, Vanilla JavaScript.\n"
            "• Backend API: Node.js, Express.js framework.\n"
            "• Database: MongoDB and Mongoose for object data modeling.\n"
            "• Authentication: JSON Web Tokens (JWT) and bcryptjs for secure passwords.\n"
            "• Third-party integrations: Razorpay for payment gateway."
        )
        for r in p.runs: format_run(r, font_size=10)

    elif text.startswith("Programming Language:"):
        p.text = "Programming Language: JavaScript (ES6+), HTML5, CSS3"
        for r in p.runs: format_run(r, font_size=10)
    elif text.startswith("Framework:"):
        p.text = "Framework: Node.js, Express.js"
        for r in p.runs: format_run(r, font_size=10)
    elif text.startswith("Database:"):
        p.text = "Database: MongoDB (Mongoose ORM)"
        for r in p.runs: format_run(r, font_size=10)
    elif text.startswith("Hardware/Software Tools:"):
        p.text = "Hardware/Software Tools: Visual Studio Code, Postman, Git, MongoDB Atlas"
        for r in p.runs: format_run(r, font_size=10)
    elif text.startswith("Other Technologies:"):
        p.text = "Other Technologies: Razorpay API, JWT, bcryptjs, CORS"
        for r in p.runs: format_run(r, font_size=10)

    elif text.startswith("Algorithm / Method followed"):
        p.text = "• Algorithm / Method followed: RESTful API architecture with secure authentication pipelines."
        for r in p.runs: format_run(r, font_size=10)
    elif text.startswith("Coding / Development details"):
        p.text = "• Coding / Development details: Developed models, routes, and controllers for products, users, and orders on the backend; connected them dynamically via frontend fetch calls."
        for r in p.runs: format_run(r, font_size=10)
    elif text.startswith("Database design"):
        p.text = "• Database design: Non-relational (NoSQL) schemas for Users, Products, Addresses, and Orders utilizing MongoDB."
        for r in p.runs: format_run(r, font_size=10)
    elif text.startswith("API integration"):
        p.text = "• API integration: Integrated Razorpay API for live checkout payment processes."
        for r in p.runs: format_run(r, font_size=10)
    elif text.startswith("Testing performed"):
        p.text = "• Testing performed: Tested API endpoints via Postman, and validated user flows like cart management and secure login on the frontend."
        for r in p.runs: format_run(r, font_size=10)

    elif text.startswith("Screenshot 1:"):
        p.text = "Screenshot 1: Frontend E-commerce UI"
        for r in p.runs: format_run(r, font_size=10, bold=True)
    elif text.startswith("Description:") and p._p.getprevious() is not None and "Screenshot 1" in p._p.getprevious().text:
        p.text = "Description: Showcases the responsive homepage and product grids."
        for r in p.runs: format_run(r, font_size=9.5, italic=True)
    elif text.startswith("Screenshot 2:"):
        p.text = "Screenshot 2: Backend API and Database"
        for r in p.runs: format_run(r, font_size=10, bold=True)
    elif text.startswith("Description:") and p._p.getprevious() is not None and "Screenshot 2" in p._p.getprevious().text:
        p.text = "Description: Demonstrates MongoDB connected via Express.js handling product fetch and payment flow."
        for r in p.runs: format_run(r, font_size=9.5, italic=True)

    elif text.startswith("Remaining modules"):
        p.text = "• Remaining modules: Admin dashboard and advanced analytics generation."
        for r in p.runs: format_run(r, font_size=10)
    elif text.startswith("Testing plan"):
        p.text = "• Testing plan: Write automated unit tests for Node.js API endpoints."
        for r in p.runs: format_run(r, font_size=10)
    elif text.startswith("Documentation work"):
        p.text = "• Documentation work: Complete README and API spec generation using Swagger."
        for r in p.runs: format_run(r, font_size=10)
    elif text.startswith("Deployment plan"):
        p.text = "• Deployment plan: Dockerize backend API and deploy on Render; host frontend on Vercel."
        for r in p.runs: format_run(r, font_size=10)

    elif text.startswith("Technical Skills Learned:"):
        p.text = "• Technical Skills Learned: Building secure backend APIs with Express, interacting with MongoDB, working with asynchronous JS, and JWT."
        for r in p.runs: format_run(r, font_size=10)
    elif text.startswith("Professional Skills Learned:"):
        p.text = "• Professional Skills Learned: System architecture planning, secure handling of user credentials, third-party API integration."
        for r in p.runs: format_run(r, font_size=10)

if len(doc.tables) > 2:
    t2 = doc.tables[2]
    try:
        t2.cell(1, 0).text = "User Auth & JWT"
        t2.cell(1, 1).text = "Implemented registration, login, and secure token generation using bcrypt and jsonwebtoken."
        t2.cell(1, 2).text = "Completed"

        modules_list = [
            ("MongoDB Models", "Created Mongoose schemas for users, products, addresses, and orders.", "Completed"),
            ("Razorpay Integration", "Setup server-side payment generation and client-side checkout.", "Completed"),
            ("Frontend UI", "Connected HTML/JS views with backend APIs via fetch to display live products and cart.", "Completed"),
            ("Cart & Checkout", "Managed cart state globally and successfully passed order requests to backend.", "Completed")
        ]

        for mod, desc, stat in modules_list:
            row = t2.add_row()
            row.cells[0].text = mod
            row.cells[1].text = desc
            row.cells[2].text = stat

        for row in t2.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        format_run(r, font_size=9.5)
    except Exception:
        pass

if len(doc.tables) > 3:
    t3 = doc.tables[3]
    try:
        t3.cell(1, 0).text = "Securing API endpoints and handling authentication states."
        t3.cell(1, 1).text = "Used Express middleware to verify JWT tokens and inject authenticated user ID into request headers."

        row_ch2 = t3.add_row()
        row_ch2.cells[0].text = "Ensuring reliable payment transactions with Razorpay API."
        row_ch2.cells[1].text = "Verified Razorpay signatures securely in the backend using crypto module to prevent tampering."

        for row in t3.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        format_run(r, font_size=9.5)
    except Exception:
        pass

for p in doc.paragraphs:
    if "References / Resources Used" in p.text or "Students must mention minimum 5 references." in p.text:
        p.text = (
            "9. References\n\n"
            "[1] Node.js Documentation, https://nodejs.org/docs\n"
            "[2] Express.js Framework Guide, https://expressjs.com/\n"
            "[3] MongoDB & Mongoose Reference, https://mongoosejs.com/docs/guide.html\n"
            "[4] Razorpay Node.js Integration, https://razorpay.com/docs/\n"
            "[5] MDN Web Docs, 'Fetch API', Mozilla Developer Network, 2026."
        )
        for r in p.runs:
            format_run(r, font_size=9.5)
        break

doc.save(output_path)
print("Updated report saved cleanly to:", output_path)
