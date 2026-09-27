import docx

template_path = r"C:\Users\krish\Downloads\Internship_UDP_Training_Report_Template for reporting 2.docx"
doc = docx.Document(template_path)

with open("template_structure.txt", "w", encoding="utf-8") as f:
    f.write("=== PARAGRAPHS ===\n")
    for i, p in enumerate(doc.paragraphs):
        f.write(f"P{i}: {p.text.strip()}\n")

    f.write("\n=== TABLES ===\n")
    for t_idx, table in enumerate(doc.tables):
        f.write(f"\n--- Table {t_idx} ({len(table.rows)} rows, {len(table.columns)} cols) ---\n")
        for r_idx, row in enumerate(table.rows):
            row_str = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
            f.write(f"Row {r_idx}: {row_str}\n")

print("Structure saved to template_structure.txt")
